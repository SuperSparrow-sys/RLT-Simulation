from datetime import datetime
from pathlib import Path

import pytest

from app import create_app
from core import anlagen, database, ergebnisse, solver
from core.wetter import einlesen, speicher, tabelle

REFERENZ = Path(__file__).parent.parent / "referenz" / "RLTSimulation_Vorlage_AX_SIM_2.1.xls"


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.DB_PATH", tmp_path / "test.db")
    anwendung = create_app()
    with anwendung.app_context():
        database.init_db()
        yield anwendung


def test_excel_datum_wird_umgerechnet():
    # 36526 ist der 1. Januar 2000 in der Excel-Zaehlung
    assert tabelle.excel_datum(36526.0).date() == datetime(2000, 1, 1).date()


def test_excel_datum_mit_uhrzeit():
    zeitpunkt = tabelle.excel_datum(36526.041666666664)
    assert zeitpunkt.hour == 1


def test_try_datei_wird_vollstaendig_gelesen():
    stunden = einlesen.lese_datei(REFERENZ)
    assert len(stunden) == 8760
    assert stunden[0]["t_au"] == pytest.approx(2.5)
    assert stunden[0]["x_au"] == pytest.approx(4.4)
    assert stunden[0]["str_h"] == pytest.approx(0.0)


def test_mittags_im_sommer_scheint_die_sonne():
    stunden = einlesen.lese_datei(REFERENZ)
    sommer = [
        s for s in stunden
        if s["zeitpunkt"].month == 7 and s["zeitpunkt"].hour == 12
    ]
    assert max(s["str_h"] for s in sommer) > 100.0


def test_datensatz_wird_gespeichert_und_gelesen(app):
    stunden = [
        {
            "zeitpunkt": datetime(2024, 1, 1, i), "t_au": float(i), "x_au": 4.0,
            "str_s": 0.0, "str_o": 0.0, "str_w": 0.0, "str_n": 0.0, "str_h": 0.0,
        }
        for i in range(24)
    ]
    with app.app_context():
        datensatz = speicher.datensatz_anlegen("Test 2024", "upload", stunden)
        zurueck = speicher.lade_stunden(datensatz)
    assert len(zurueck) == 24
    assert zurueck[5]["t_au"] == pytest.approx(5.0)
    assert zurueck[0]["zeitpunkt"].hour == 0


def test_ausschnitt_laesst_sich_laden(app):
    stunden = [
        {
            "zeitpunkt": datetime(2024, 1, 1, i % 24), "t_au": float(i), "x_au": 4.0,
            "str_s": 0.0, "str_o": 0.0, "str_w": 0.0, "str_n": 0.0, "str_h": 0.0,
        }
        for i in range(100)
    ]
    with app.app_context():
        datensatz = speicher.datensatz_anlegen("Test", "upload", stunden)
        ausschnitt = speicher.lade_stunden(datensatz, von=10, bis=20)
    assert len(ausschnitt) == 10
    assert ausschnitt[0]["t_au"] == pytest.approx(10.0)


def _schreibe_probe_xlsx(pfad):
    """Eine kleine Mappe, deren Datumsspalte als Datum formatiert ist.

    Genau so sieht die Wetterdatei aus, wenn jemand die Vorlage in Excel oeffnet
    und als .xlsx speichert: openpyxl liefert dann fertige Zeitstempel statt
    Tageszahlen.
    """
    import openpyxl

    mappe = openpyxl.Workbook()
    blatt = mappe.active
    blatt.title = "Wetterdaten"
    for zeile in range(1, 5):
        blatt.cell(zeile, 1, "Kopfzeile")
    for i in range(3):
        blatt.cell(5 + i, 1, datetime(2000, 1, 1, 1 + i))
        for spalte, wert in enumerate([2.5 + i, 4.4, 0.0, 0.0, 0.0, 0.0, 0.0], start=2):
            blatt.cell(5 + i, spalte, wert)
    mappe.save(pfad)


def test_xlsx_mit_datumsformatierten_zellen(tmp_path):
    """Sonst liest das Programm aus einer gespeicherten Mappe null Stunden."""
    pfad = tmp_path / "probe.xlsx"
    _schreibe_probe_xlsx(pfad)
    stunden = einlesen.lese_datei(pfad)

    assert len(stunden) == 3
    assert stunden[0]["zeitpunkt"] == datetime(2000, 1, 1, 1)
    assert stunden[0]["t_au"] == pytest.approx(2.5)
    assert stunden[2]["t_au"] == pytest.approx(4.5)


def test_csv_mit_beiden_trennzeichen(tmp_path):
    for trenner in (",", ";"):
        pfad = tmp_path / f"probe{'komma' if trenner == ',' else 'semikolon'}.csv"
        zeilen = ["Kopf"] * 4 + [
            trenner.join(["36526.041666666664", "2.5", "4.4", "0", "0", "0", "0", "0"]),
            trenner.join(["36526.083333333336", "3.1", "4.6", "0", "0", "0", "0", "0"]),
        ]
        pfad.write_text("\n".join(zeilen), encoding="utf-8")
        stunden = einlesen.lese_datei(pfad)
        assert len(stunden) == 2, trenner
        assert stunden[0]["t_au"] == pytest.approx(2.5), trenner
        assert stunden[1]["zeitpunkt"].hour == 2, trenner


def test_upload_meldet_eine_kaputte_datei_lesbar(app):
    """Eine beschaedigte Datei muss einen Satz ergeben, keine Fehlerseite."""
    import io

    antwort = app.test_client().post(
        "/api/wetter/upload",
        data={"datei": (io.BytesIO(b"kein Tabellendokument"), "kaputt.xls")},
        content_type="multipart/form-data",
    )
    assert antwort.status_code == 400
    assert "lesen" in antwort.get_json()["fehler"]


def test_upload_nimmt_die_mappe_an(app):
    antwort = app.test_client().post(
        "/api/wetter/upload",
        data={"datei": (open(REFERENZ, "rb"), "TRY04.xls"), "name": "TRY04"},
        content_type="multipart/form-data",
    )
    assert antwort.status_code == 201
    assert antwort.get_json()["stunden"] == 8760


def test_api_listet_die_datensaetze(app):
    with app.app_context():
        speicher.datensatz_anlegen(
            "TRY04", "upload",
            [
                {
                    "zeitpunkt": datetime(2024, 1, 1, 0), "t_au": 1.0, "x_au": 4.0,
                    "str_s": 0.0, "str_o": 0.0, "str_w": 0.0, "str_n": 0.0, "str_h": 0.0,
                }
            ],
        )
    antwort = app.test_client().get("/api/wetter")
    assert antwort.status_code == 200
    assert antwort.get_json()[0]["name"] == "TRY04"


# -- Loeschen und Umbenennen -----------------------------------------------

def _stunden(anzahl=1):
    return [
        {
            "zeitpunkt": datetime(2024, 1, 1, i % 24), "t_au": 0.0, "x_au": 4.0,
            "str_s": 0.0, "str_o": 0.0, "str_w": 0.0, "str_n": 0.0, "str_h": 0.0,
        }
        for i in range(anzahl)
    ]


def test_datensatz_umbenennen(app):
    with app.app_context():
        datensatz = speicher.datensatz_anlegen("Alt", "upload", _stunden())
        speicher.datensatz_umbenennen(datensatz, "Neu")
        name = database.get_db().execute(
            "SELECT name FROM wetterdatensatz WHERE id = ?", (datensatz,)
        ).fetchone()["name"]
    assert name == "Neu"


def test_datensatz_umbenennen_unbekannt_meldet_fehler(app):
    with app.app_context():
        with pytest.raises(KeyError):
            speicher.datensatz_umbenennen(9999, "Neu")


def test_datensatz_loeschen_entfernt_seine_stunden(app):
    with app.app_context():
        datensatz = speicher.datensatz_anlegen("W", "upload", _stunden(5))
        speicher.datensatz_loeschen(datensatz)
        db = database.get_db()
        assert db.execute(
            "SELECT COUNT(*) AS n FROM wetterdatensatz WHERE id = ?", (datensatz,)
        ).fetchone()["n"] == 0
        assert db.execute(
            "SELECT COUNT(*) AS n FROM wetterstunde WHERE datensatz_id = ?", (datensatz,)
        ).fetchone()["n"] == 0


def test_datensatz_loeschen_wird_verweigert_wenn_ein_lauf_darauf_verweist(app):
    """Ein gespeicherter Simulationslauf ohne seinen Wetterdatensatz waere
    nicht mehr nachvollziehbar - siehe core/wetter/speicher.py."""
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage = anlagen.anlage_anlegen(projekt, "A")
        datensatz = speicher.datensatz_anlegen("W", "upload", _stunden())
        graph = anlagen.lade_graph(anlage)
        lauf = solver.Lauf(
            stunden=[{}], bilanz={"strom_ht": 0.0, "strom_nt": 0.0, "waerme": 0.0,
                                   "kaelte": 0.0, "wasser": 0.0},
            warnungen=[],
        )
        ergebnisse.speichere(anlage, datensatz, 0, 1, lauf, graph, dauer=0.1)

        with pytest.raises(ValueError, match="wird von 1 Simulationslauf verwendet"):
            speicher.datensatz_loeschen(datensatz)

        # Weiterhin vollstaendig vorhanden - kein Teilloeschen.
        db = database.get_db()
        assert db.execute(
            "SELECT COUNT(*) AS n FROM wetterdatensatz WHERE id = ?", (datensatz,)
        ).fetchone()["n"] == 1


def test_api_datensatz_umbenennen_und_loeschen(app):
    klient = app.test_client()
    with app.app_context():
        datensatz = speicher.datensatz_anlegen("Alt", "upload", _stunden())

    antwort = klient.patch(f"/api/wetter/{datensatz}", json={"name": "Neu"})
    assert antwort.status_code == 200
    assert klient.get("/api/wetter").get_json()[0]["name"] == "Neu"

    antwort = klient.delete(f"/api/wetter/{datensatz}")
    assert antwort.status_code == 200
    assert klient.get("/api/wetter").get_json() == []


def test_api_datensatz_loeschen_verweigert_bei_verweisendem_lauf(app):
    klient = app.test_client()
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage = anlagen.anlage_anlegen(projekt, "A")
        datensatz = speicher.datensatz_anlegen("W", "upload", _stunden())
        graph = anlagen.lade_graph(anlage)
        lauf = solver.Lauf(
            stunden=[{}], bilanz={"strom_ht": 0.0, "strom_nt": 0.0, "waerme": 0.0,
                                   "kaelte": 0.0, "wasser": 0.0},
            warnungen=[],
        )
        ergebnisse.speichere(anlage, datensatz, 0, 1, lauf, graph, dauer=0.1)

    antwort = klient.delete(f"/api/wetter/{datensatz}")
    assert antwort.status_code == 400
    assert "Simulationslauf" in antwort.get_json()["fehler"]
