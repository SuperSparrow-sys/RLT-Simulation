"""core/ausgabe.py und die Ausgabe-Routen in routes/simulation.py.

Baut Testlaeufe wie tests/test_bericht.py: ein synthetischer core.solver.Lauf
statt eines echten (minutenlangen) Solver-Durchlaufs - core.ergebnisse.speichere()
speichert ihn genauso wie einen echten. Anders als test_bericht.py fuellt der
Lauf hier zusaetzlich die Datenlogger-Karte der Vorlage (ax_sim_2_1 hat vier
benannte Spalten), damit sich auch das Stundenprotokoll in der Ausgabe pruefen
laesst.
"""

import csv
import io
import re
import time
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from app import create_app
from core import anlagen, ausgabe, database, ergebnisse, solver
from core.vorlagen import ax_sim_2_1
from core.wetter import speicher


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr("core.config.DB_PATH", tmp_path / "test.db")
    anwendung = create_app()
    with anwendung.app_context():
        database.init_db()
        yield anwendung


def _wetter_anlegen(anzahl, jahr=2023):
    start = datetime(jahr, 1, 1)
    return speicher.datensatz_anlegen(
        "Testwetter", "upload",
        [
            {
                "zeitpunkt": start + timedelta(hours=i),
                "t_au": -5.0 + i, "x_au": 3.0 + i * 0.1,
                "str_s": 0.0, "str_o": 0.0, "str_w": 0.0, "str_n": 0.0, "str_h": 0.0,
            }
            for i in range(anzahl)
        ],
        ort="Musterstadt", jahr=jahr,
    )


def _karte_id(graph, typ):
    return next(k.id for k in graph.karten.values() if k.typ == typ)


def _lauf_speichern(app, anzahl_stunden=5, anlage_name="A", jahr=2023):
    """Ein kurzer, synthetischer Lauf, der sowohl die Bilanzkarte als auch die
    Datenlogger-Karte der Vorlage bestueckt (vier benannte Spalten:
    'WRG'/'T_ZU WRG'/'T Raum'/'F Raum', siehe core/vorlagen/ax_sim_2_1.py)."""
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage_id = ax_sim_2_1.baue(projekt, anlage_name)
        wetter_id = _wetter_anlegen(anzahl_stunden, jahr=jahr)
        graph = anlagen.lade_graph(anlage_id)
        bilanz_id = _karte_id(graph, "bilanz")
        logger_id = _karte_id(graph, "datenlogger")

        stunden = [
            {
                bilanz_id: {
                    "waerme": float(10 + n), "kaelte": float(n % 3),
                    "strom_ht": 1.0 + n, "strom_nt": 0.5, "wasser": 0.2 * n,
                },
                logger_id: {
                    "wert_1": float(n), "wert_2": float(n * 2),
                    "wert_3": 20.0 + n, "wert_4": 5.0 + n * 0.5,
                },
            }
            for n in range(anzahl_stunden)
        ]
        lauf = solver.Lauf(
            stunden=stunden,
            bilanz={
                "waerme": sum(s[bilanz_id]["waerme"] for s in stunden),
                "kaelte": sum(s[bilanz_id]["kaelte"] for s in stunden),
                "strom_ht": sum(s[bilanz_id]["strom_ht"] for s in stunden),
                "strom_nt": anzahl_stunden * 0.5,
                "wasser": sum(s[bilanz_id]["wasser"] for s in stunden),
            },
            warnungen=[],
        )
        simulation_id = ergebnisse.speichere(
            anlage_id, wetter_id, 0, anzahl_stunden, lauf, graph, dauer=0.2
        )
    return anlage_id, simulation_id


# ---------------------------------------------------------------------------
# core.ausgabe.daten_fuer()
# ---------------------------------------------------------------------------

def test_daten_fuer_unbekannte_simulation_wirft_keyerror(app):
    with app.app_context():
        with pytest.raises(KeyError):
            ausgabe.daten_fuer(999999)


def test_daten_fuer_laufender_lauf_wirft_nicht_verfuegbar(app):
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage_id = ax_sim_2_1.baue(projekt, "A")
        wetter_id = _wetter_anlegen(3)
        simulation_id = ergebnisse.beginne(anlage_id, wetter_id, 0, 3, "kennung-1")
        with pytest.raises(ausgabe.AusgabeNichtVerfuegbar):
            ausgabe.daten_fuer(simulation_id)


def test_daten_fuer_enthaelt_alle_erwarteten_spalten(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=5)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)

    namen = [s["name"] for s in daten["spalten"]]
    # Zeitpunkt, Außentemperatur, Außenfeuchte
    assert namen[:3] == ["Zeitpunkt", "Außentemperatur", "Außenfeuchte"]
    # die fuenf Bilanzgroessen
    for label in ["Strom Hochtarif", "Strom Niedertarif", "Wärme", "Kälte", "Wasser"]:
        assert label in namen
    # alle vier benannten Datenlogger-Spalten der Vorlage
    for logger_name in ["WRG", "T_ZU WRG", "T Raum", "F Raum"]:
        assert logger_name in namen

    # jede Spalte hat einen Wert je gerechneter Stunde
    for spalte in daten["spalten"]:
        assert len(spalte["werte"]) == 5


def test_daten_fuer_werte_stimmen_mit_wetter_und_bilanz_ueberein(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=4)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
    spalten = {s["name"]: s["werte"] for s in daten["spalten"]}

    assert [z.hour for z in spalten["Zeitpunkt"]] == [0, 1, 2, 3]
    assert spalten["Außentemperatur"] == pytest.approx([-5.0, -4.0, -3.0, -2.0])
    assert spalten["Wärme"] == pytest.approx([10.0, 11.0, 12.0, 13.0])
    assert spalten["WRG"] == pytest.approx([0.0, 1.0, 2.0, 3.0])
    assert spalten["T Raum"] == pytest.approx([20.0, 21.0, 22.0, 23.0])


def test_daten_fuer_ohne_bilanzkarte_hat_trotzdem_wetterspalten(app):
    """Eine Anlage ganz ohne Bilanzkarte darf die Ausgabe nicht mit einem
    Fehler abbrechen lassen - nur eben ohne Bilanzspalten (wie
    core.bericht.daten_fuer() fuer die Diagramme)."""
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage_id = anlagen.anlage_anlegen(projekt, "Leer")
        wetter_id = _wetter_anlegen(3)
        graph = anlagen.lade_graph(anlage_id)
        lauf = solver.Lauf(stunden=[{}, {}, {}], bilanz={}, warnungen=[])
        simulation_id = ergebnisse.speichere(anlage_id, wetter_id, 0, 3, lauf, graph, dauer=0.1)
        daten = ausgabe.daten_fuer(simulation_id)
    namen = [s["name"] for s in daten["spalten"]]
    assert namen == ["Zeitpunkt", "Außentemperatur", "Außenfeuchte"]


def test_daten_fuer_abgebrochener_lauf_ist_verfuegbar(app):
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage_id = ax_sim_2_1.baue(projekt, "A")
        wetter_id = _wetter_anlegen(5)
        graph = anlagen.lade_graph(anlage_id)
        bilanz_id = _karte_id(graph, "bilanz")
        # Nur 2 von 5 angeforderten Stunden gerechnet - wie ein Abbruch.
        stunden = [{bilanz_id: {"waerme": 1.0}} for _ in range(2)]
        lauf = solver.Lauf(stunden=stunden, bilanz={"waerme": 2.0}, warnungen=[])
        simulation_id = ergebnisse.speichere(
            anlage_id, wetter_id, 0, 5, lauf, graph, dauer=0.1, status="abgebrochen"
        )
        daten = ausgabe.daten_fuer(simulation_id)
    spalten = {s["name"]: s["werte"] for s in daten["spalten"]}
    assert len(spalten["Zeitpunkt"]) == 2  # nicht die angeforderten 5


# ---------------------------------------------------------------------------
# core.ausgabe.dateiname()
# ---------------------------------------------------------------------------

def test_dateiname_nennt_anlage_wetterjahr_und_laufdatum(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=2, anlage_name="Büro-Nord", jahr=2021)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        name = ausgabe.dateiname(daten, "csv")
    assert name.startswith("stundenwerte-")
    assert "2021" in name
    assert name.endswith(".csv")
    assert "Nord" in name  # der Bindestrich im Anlagennamen bleibt erhalten


def test_dateiname_ohne_sonderzeichen(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=2, anlage_name='Süd "Test"/Halle')
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        name = ausgabe.dateiname(daten, "xlsx")
    assert '"' not in name and "/" not in name


# ---------------------------------------------------------------------------
# core.ausgabe.csv_bytes() - Semikolon, Komma als Dezimalzeichen, UTF-8-BOM
# ---------------------------------------------------------------------------

def test_csv_bytes_hat_utf8_bom(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=3)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.csv_bytes(daten)
    assert rohdaten.startswith(b"\xef\xbb\xbf")


def test_csv_bytes_oeffnet_mit_deutscher_excel_einstellung(app):
    """Simuliert, was eine deutsche Excel-Einstellung tut: BOM erkennen,
    Semikolon als Feldtrennzeichen, Komma als Dezimalzeichen lesen."""
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=4)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.csv_bytes(daten)

    text = rohdaten.decode("utf-8-sig")
    zeilen = list(csv.reader(io.StringIO(text), delimiter=";"))
    kopf = zeilen[0]
    assert kopf[0] == "Zeitpunkt"
    assert "Außentemperatur [°C]" in kopf
    assert "Wärme [kWh]" in kopf
    assert "WRG" in " ".join(kopf)  # Datenlogger-Spalte ohne Einheit -> kein '['

    assert len(zeilen) == 1 + 4  # Kopfzeile + vier Stunden

    temp_index = kopf.index("Außentemperatur [°C]")
    erste_datenzeile = zeilen[1]
    # Punkt darf im Zahlenfeld nicht vorkommen (Verwechslung mit Dezimalpunkt);
    # das deutsche Dezimalkomma steht stattdessen dort.
    assert "." not in erste_datenzeile[temp_index]
    assert "," in erste_datenzeile[temp_index]
    wert = float(erste_datenzeile[temp_index].replace(",", "."))
    assert wert == pytest.approx(-5.0)

    # Semikolon taucht innerhalb keines Zahlenfelds auf - die Zeile hat pro
    # Spalte genau ein Feld, csv.reader hat also richtig getrennt.
    assert len(erste_datenzeile) == len(kopf)


def test_csv_bytes_zeitpunkt_ist_deutsches_datumsformat(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=2)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.csv_bytes(daten)
    text = rohdaten.decode("utf-8-sig")
    zeilen = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert re.match(r"^\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}$", zeilen[1][0])


# ---------------------------------------------------------------------------
# core.ausgabe.xlsx_bytes() - echte Zahlen, Kopfzeile, Spaltenbreiten, Fixierung
# ---------------------------------------------------------------------------

def test_xlsx_bytes_lesbar_mit_richtigen_massen(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=6)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.xlsx_bytes(daten)

    arbeitsmappe = load_workbook(io.BytesIO(rohdaten))
    blatt = arbeitsmappe["Stundenwerte"]

    anzahl_spalten = len(daten["spalten"])
    assert blatt.max_row == 1 + 6
    assert blatt.max_column == anzahl_spalten
    assert blatt.freeze_panes == "A2"

    kopfzeile = [z.value for z in blatt[1]]
    assert kopfzeile[0] == "Zeitpunkt"
    assert any("Wärme" in str(w) for w in kopfzeile)
    assert blatt.cell(row=1, column=1).font.bold is True

    for spalte in range(1, anzahl_spalten + 1):
        assert blatt.column_dimensions[blatt.cell(1, spalte).column_letter].width


def test_xlsx_bytes_zahlen_sind_zahlen_nicht_text(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=3)
    with app.app_context():
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.xlsx_bytes(daten)
    blatt = load_workbook(io.BytesIO(rohdaten))["Stundenwerte"]

    kopfzeile = [z.value for z in blatt[1]]
    temp_spalte = kopfzeile.index("Außentemperatur [°C]") + 1
    zeitpunkt_spalte = 1

    zweite_zeile = blatt[2]
    # openpyxl liest eine ganzzahlige Zahl beim Zuruecklesen als int statt
    # float (Excel selbst unterscheidet intern nicht zwischen beidem) - fuer
    # 'Zahl statt Text' zaehlt nur, dass es KEIN str ist.
    wert = zweite_zeile[temp_spalte - 1].value
    assert isinstance(wert, (int, float)) and not isinstance(wert, bool)
    assert wert == pytest.approx(-5.0)
    assert isinstance(zweite_zeile[zeitpunkt_spalte - 1].value, datetime)


def test_xlsx_bytes_stundenzahl_und_spaltenwerte_stimmen_mit_protokoll_ueberein(app):
    """Was das Stundenprotokoll im Dialog zeigt (core.ergebnisse.lade_protokoll)
    muss in der Excel-Ausgabe wiederzufinden sein - dieselben Spaltennamen und
    dieselben Werte."""
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=5)
    with app.app_context():
        graph = anlagen.lade_graph(_anlage_id)
        protokoll = ergebnisse.lade_protokoll(simulation_id, graph)
        daten = ausgabe.daten_fuer(simulation_id)
        rohdaten = ausgabe.xlsx_bytes(daten)

    blatt = load_workbook(io.BytesIO(rohdaten))["Stundenwerte"]
    kopfzeile = [z.value for z in blatt[1]]

    for spalte in protokoll:
        gesuchter_name = f"{spalte['name']} [{spalte['einheit']}]" if spalte["einheit"] else spalte["name"]
        assert gesuchter_name in kopfzeile
        index = kopfzeile.index(gesuchter_name)
        excel_werte = [blatt.cell(row=r, column=index + 1).value for r in range(2, 2 + 5)]
        assert excel_werte == pytest.approx(spalte["werte"], rel=1e-4)


# ---------------------------------------------------------------------------
# routes/simulation.py - /stundenwerte.csv und /stundenwerte.xlsx
# ---------------------------------------------------------------------------

def test_route_csv_liefert_anhang(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=3)
    antwort = app.test_client().get(f"/api/simulation/{simulation_id}/stundenwerte.csv")
    assert antwort.status_code == 200
    assert antwort.headers["Content-Type"].startswith("text/csv")
    assert "attachment" in antwort.headers["Content-Disposition"]
    assert antwort.data.startswith(b"\xef\xbb\xbf")


def test_route_xlsx_liefert_anhang(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=3)
    antwort = app.test_client().get(f"/api/simulation/{simulation_id}/stundenwerte.xlsx")
    assert antwort.status_code == 200
    assert antwort.headers["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in antwort.headers["Content-Disposition"]
    # Eine .xlsx-Datei ist ein Zip-Archiv - beginnt mit der Zip-Magic 'PK'.
    assert antwort.data[:2] == b"PK"


def test_route_dateiname_ohne_sonderzeichen(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=2, anlage_name='Halle "1"/Süd')
    antwort = app.test_client().get(f"/api/simulation/{simulation_id}/stundenwerte.csv")
    assert antwort.status_code == 200
    name = re.search(r'filename="([^"]+)"', antwort.headers["Content-Disposition"]).group(1)
    assert '"' not in name and "/" not in name


def test_route_unbekannte_simulation_ist_404(app):
    antwort = app.test_client().get("/api/simulation/999999/stundenwerte.csv")
    assert antwort.status_code == 404
    antwort = app.test_client().get("/api/simulation/999999/stundenwerte.xlsx")
    assert antwort.status_code == 404


def test_route_laufender_lauf_ist_409(app):
    with app.app_context():
        projekt = anlagen.projekt_anlegen("P")
        anlage_id = ax_sim_2_1.baue(projekt, "A")
        wetter_id = _wetter_anlegen(5)
        simulation_id = ergebnisse.beginne(anlage_id, wetter_id, 0, 5, "kennung-2")
    antwort = app.test_client().get(f"/api/simulation/{simulation_id}/stundenwerte.csv")
    assert antwort.status_code == 409
    antwort = app.test_client().get(f"/api/simulation/{simulation_id}/stundenwerte.xlsx")
    assert antwort.status_code == 409


# ---------------------------------------------------------------------------
# Laufzeit bei einem vollen Jahreslauf (8760 Zeilen) - kuenstlich erzeugt statt
# tatsaechlich gerechnet (acht Minuten Solver-Laufzeit sind hierfuer nicht
# noetig), damit die Ausgabe nicht laenger als eine schlichte Anfrage dauern
# darf. Als 'slow' markiert wie die anderen laenger laufenden Pruefungen
# (siehe tests/test_abgleich.py) - fuer die uebliche Testreihe uninteressant,
# aber bei einer Aenderung an core.ausgabe hier erneut zu pruefen.
# ---------------------------------------------------------------------------

@pytest.mark.slow
def test_ausgabe_bei_8760_zeilen_ist_schnell_genug(app):
    _anlage_id, simulation_id = _lauf_speichern(app, anzahl_stunden=8760)
    with app.app_context():
        begonnen = time.time()
        daten = ausgabe.daten_fuer(simulation_id)
        csv_rohdaten = ausgabe.csv_bytes(daten)
        xlsx_rohdaten = ausgabe.xlsx_bytes(daten)
        dauer = time.time() - begonnen

    assert len(daten["spalten"][0]["werte"]) == 8760
    assert len(csv_rohdaten) > 0
    assert len(xlsx_rohdaten) > 0
    # Grosszuegige Grenze: auf der Entwicklungsmaschine lagen beide Formate
    # zusammen bei rund 2 Sekunden (siehe ausgabe-report.md) - ausreichend
    # schnell fuer eine schlichte Anfrage ohne Hintergrundmuster.
    assert dauer < 10.0
