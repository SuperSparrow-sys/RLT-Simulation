"""Liest Wetterdaten im Format des Excel-Blattes 'Wetterdaten'.

Aufbau dort: Zeile 1 bis 4 Kopf, ab Zeile 5 je Stunde
Datum als Excel-Zahl, Temperatur in °C, absolute Feuchte in g/kg und
die Strahlung auf Sued, Ost, West, Nord und die Horizontale in W/m².

Zustaendig fuer .xls, .xlsx und .csv. Ein DWD-Testreferenzjahr im .dat-Format
liest dieses Modul nicht - das tut core/wetter/try_dat.py. Bis diese Trennung
entstand, hiess die Datei hier "try_import.py" und las trotzdem nur Excel; der
Name legte eine Faehigkeit nahe, die das Programm gar nicht hatte.
"""

import csv
from datetime import date, datetime, timedelta
from pathlib import Path

SPALTEN = ("t_au", "x_au", "str_s", "str_o", "str_w", "str_n", "str_h")

# Excel zaehlt Tage ab dem 30.12.1899 (mit dem bekannten Schaltjahrfehler von 1900,
# der fuer Datumsangaben ab 1900 keine Rolle mehr spielt).
EXCEL_NULLPUNKT = datetime(1899, 12, 30)


def excel_datum(wert, jahr=None):
    """Macht aus einer Datumsangabe einen Zeitstempel auf voller Stunde.

    Der Wert kommt je nach Dateiart unterschiedlich an: aus einer .xls-Datei und
    aus CSV als Tageszahl seit dem 30.12.1899, aus einer .xlsx-Datei dagegen als
    fertiger Zeitstempel, weil openpyxl datumsformatierte Zellen selbst umrechnet.
    Beide Formen muessen hier durch - sonst liest das Programm aus einer als .xlsx
    gespeicherten Mappe keine einzige Stunde und meldet nur, die Datei enthalte
    keine Werte.
    """
    if isinstance(wert, datetime):
        zeitpunkt = wert
    elif isinstance(wert, date):
        zeitpunkt = datetime(wert.year, wert.month, wert.day)
    else:
        zeitpunkt = EXCEL_NULLPUNKT + timedelta(days=float(wert))

    # Auf volle Stunden runden - die Excel speichert 0,0416666666 statt 1/24
    zeitpunkt += timedelta(seconds=30 * 60)
    zeitpunkt = zeitpunkt.replace(minute=0, second=0, microsecond=0)
    if jahr is not None and not (zeitpunkt.month == 2 and zeitpunkt.day == 29):
        zeitpunkt = zeitpunkt.replace(year=jahr)
    return zeitpunkt


def _zeilen_aus_xls(pfad):
    import xlrd

    mappe = xlrd.open_workbook(str(pfad))
    blatt = (
        mappe.sheet_by_name("Wetterdaten")
        if "Wetterdaten" in mappe.sheet_names()
        else mappe.sheet_by_index(0)
    )
    for nummer in range(4, blatt.nrows):
        yield blatt.row_values(nummer)[:8]


def _zeilen_aus_xlsx(pfad):
    import openpyxl

    mappe = openpyxl.load_workbook(str(pfad), data_only=True)
    blatt = mappe["Wetterdaten"] if "Wetterdaten" in mappe.sheetnames else mappe.worksheets[0]
    for nummer, zeile in enumerate(blatt.iter_rows(values_only=True), start=1):
        if nummer <= 4:
            continue
        yield list(zeile)[:8]


def _zeilen_aus_csv(pfad):
    with open(pfad, encoding="utf-8-sig", newline="") as datei:
        leser = csv.reader(datei, delimiter=";" if _semikolon(pfad) else ",")
        for nummer, zeile in enumerate(leser, start=1):
            if nummer <= 4:
                continue
            yield zeile[:8]


def _semikolon(pfad):
    # Nicht nur die erste Zeile pruefen: Kopfzeilen sind oft reiner Text ohne
    # Trennzeichen (z. B. "Kopfzeile"), dann liesse sich daraus kein Trenner
    # ablesen und es wuerde faelschlich immer Komma gewaehlt, selbst wenn die
    # Datenzeilen durchgehend mit Semikolon getrennt sind.
    with open(pfad, encoding="utf-8-sig") as datei:
        inhalt = datei.read()
    return inhalt.count(";") > inhalt.count(",")


def lese_datei(pfad, jahr=None):
    pfad = Path(pfad)
    endung = pfad.suffix.lower()
    if endung == ".xls":
        zeilen = _zeilen_aus_xls(pfad)
    elif endung in (".xlsx", ".xlsm"):
        zeilen = _zeilen_aus_xlsx(pfad)
    elif endung == ".csv":
        zeilen = _zeilen_aus_csv(pfad)
    else:
        raise ValueError(f"Format '{endung}' wird nicht unterstützt")

    stunden = []
    for zeile in zeilen:
        if not zeile or zeile[0] in ("", None):
            continue
        try:
            zeitpunkt = excel_datum(zeile[0], jahr)
        except (TypeError, ValueError):
            continue
        werte = {"zeitpunkt": zeitpunkt}
        for index, name in enumerate(SPALTEN, start=1):
            wert = zeile[index] if index < len(zeile) else 0.0
            werte[name] = float(wert) if wert not in ("", None) else 0.0
        stunden.append(werte)
    return stunden
